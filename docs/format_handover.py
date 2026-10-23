# -*- coding: utf-8 -*-
"""Post-process the pandoc-generated handover DOCX for Chinese report styling."""
import pathlib
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PATH = pathlib.Path(r"D:\cxdownload\javaee课设\项目交接文档.docx")
SRC = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else PATH
OUT = pathlib.Path(sys.argv[2]) if len(sys.argv) > 2 else SRC

doc = Document(SRC)

BLUE = RGBColor(0x1F, 0x38, 0x64)
GRAY = RGBColor(0x59, 0x59, 0x59)
DARK = RGBColor(0x22, 0x22, 0x22)


def set_font(style, cn="宋体", en="Times New Roman", size=10.5, bold=None, color=None):
    style.font.name = en
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:eastAsia"), cn)


# page setup
for sec in doc.sections:
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.2)

# base styles
for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
    st = doc.styles[name]
    set_font(st, "宋体", "Times New Roman", 10.5, color=DARK)
    pf = st.paragraph_format
    pf.line_spacing = 1.45
    pf.space_before = Pt(0)
    pf.space_after = Pt(4)
    pf.first_line_indent = Pt(0)

# headings
for name, size, before, after in (
    ("Heading 1", 18, 16, 8),
    ("Heading 2", 15, 14, 6),
    ("Heading 3", 12.5, 10, 4),
    ("Heading 4", 11, 8, 3),
):
    st = doc.styles[name]
    set_font(st, "黑体", "Times New Roman", size, bold=True, color=BLUE)
    pf = st.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.keep_with_next = True

# title block
for name, size, color, before, after in (
    ("Title", 24, BLUE, 0, 6),
    ("Subtitle", 13, GRAY, 0, 10),
    ("Author", 10.5, GRAY, 0, 2),
    ("Date", 10.5, GRAY, 0, 16),
):
    st = doc.styles[name]
    set_font(st, "黑体" if name == "Title" else "宋体", "Times New Roman", size,
             bold=(name == "Title"), color=color)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)

# block quote
try:
    st = doc.styles["Block Text"]
    set_font(st, "宋体", "Times New Roman", 9.5, color=GRAY)
    st.paragraph_format.left_indent = Cm(0.4)
    st.paragraph_format.space_after = Pt(4)
except KeyError:
    pass

# source code
try:
    st = doc.styles["Source Code"]
    set_font(st, "Consolas", "Consolas", 8.5, color=RGBColor(0x1A, 0x1A, 0x1A))
    st.paragraph_format.line_spacing = 1.15
    st.paragraph_format.space_after = Pt(2)
    st.paragraph_format.left_indent = Cm(0.3)
except KeyError:
    pass


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def set_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "BFBFBF")
        borders.append(el)
    tblPr.append(borders)


for table in doc.tables:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_borders(table)
    if table.rows:
        trPr = table.rows[0]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:tblHeader"))
    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if ri == 0:
                shade(cell, "DCE6F1")
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.15
                for r in p.runs:
                    r.font.size = Pt(8.5)
                    r.font.name = "Times New Roman"
                    r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
                    if ri == 0:
                        r.bold = True

# page number footer
section = doc.sections[0]
footer_p = section.footer.paragraphs[0]
for child in list(footer_p._p):
    footer_p._p.remove(child)
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer_p.add_run()
f1 = OxmlElement("w:fldChar")
f1.set(qn("w:fldCharType"), "begin")
it = OxmlElement("w:instrText")
it.set(qn("xml:space"), "preserve")
it.text = "PAGE"
f2 = OxmlElement("w:fldChar")
f2.set(qn("w:fldCharType"), "end")
run._r.append(f1)
run._r.append(it)
run._r.append(f2)
run.font.size = Pt(9)
run.font.color.rgb = GRAY

doc.save(OUT)
print("saved", OUT)
